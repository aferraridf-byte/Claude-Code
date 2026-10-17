#!/usr/bin/env python3
"""
Relatório semanal de desempenho de conteúdo no LinkedIn.

Lê o export de análises de conteúdo do LinkedIn (.xlsx "Content_..." ou
"AggregateAnalytics_...", últimos 365 dias) e, se houver, a planilha de
detalhes por post coletada no navegador (posts_detalhe.csv). Gera:

  - Relatorio_LinkedIn_AAAA-MM-DD.xlsx  (resumo, semana x histórico, posts,
                                          padrões dos melhores posts, série semanal)
  - Relatorio_LinkedIn_AAAA-MM-DD.md    (o mesmo em texto, para leitura rápida)

e imprime um JSON com os números e os aprendizados candidatos.

Uso:
    python analisar_linkedin.py --export EXPORT.xlsx --saida "LinkedIn/Relatórios"
    python analisar_linkedin.py --export "LinkedIn/Exportações" \
        --detalhes "LinkedIn/Dados/posts_detalhe.csv" --saida "LinkedIn/Relatórios" \
        --aprendizados aprendizados.txt

O que o export do LinkedIn traz (e o que não traz):
  - ENGAGEMENT: impressões e engajamentos por dia.
  - FOLLOWERS: novos seguidores por dia e o total no último dia.
  - TOP POSTS: até 50 posts por engajamentos e até 50 por impressões (link, data e
    o número). Quem publica mais de ~50 posts no ano perde os mais fracos.
  - NÃO traz texto, formato, comentários, compartilhamentos nem seguidores por post.
    Isso vem do posts_detalhe.csv (coleta no navegador). Sem ele, o relatório sai
    só com impressões, engajamento, dia e horário.
  - Dia e horário exatos de cada post saem do ID no link: os 41 bits altos do ID
    são o instante da publicação em milissegundos (UTC).
"""
import argparse
import csv
import json
import re
import statistics
import sys
import unicodedata
from bisect import bisect_left, bisect_right
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("ERRO: falta o openpyxl. Instale com: pip install openpyxl")

DIAS = ["segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo"]
FAIXAS = [(0, 6, "0–6h"), (6, 9, "6–9h"), (9, 12, "9–12h"), (12, 14, "12–14h"),
          (14, 18, "14–18h"), (18, 24, "18–24h")]
TAMANHOS = [(500, "curto (até 500 caracteres)"), (1300, "médio (501–1.300)"),
            (10 ** 9, "longo (mais de 1.300)")]
DIMENSOES = [("formato", "Formato"), ("tipo_gancho", "Gancho"), ("tamanho", "Tamanho"),
             ("tema", "Tema"), ("dia_semana", "Dia da semana"), ("faixa_horaria", "Horário")]
METRICAS_SCORE = ["impressoes", "taxa_engajamento", "comentarios", "compartilhamentos",
                  "seguidores_ganhos"]
METRICAS_POST = ["impressoes", "engajamentos", "taxa_engajamento", "comentarios",
                 "compartilhamentos", "seguidores_ganhos"]
METRICAS_DETALHE = ["impressoes", "alcance", "reacoes", "comentarios", "compartilhamentos",
                    "salvamentos", "envios", "seguidores_ganhos", "visitas_perfil"]
ROTULOS = {
    "impressoes": "Impressões", "engajamentos": "Engajamentos",
    "taxa_engajamento": "Taxa de engajamento", "comentarios": "Comentários",
    "compartilhamentos": "Compartilhamentos", "seguidores_ganhos": "Seguidores ganhos",
    "reacoes": "Reações", "salvamentos": "Salvamentos", "envios": "Envios",
    "visitas_perfil": "Visitas ao perfil", "alcance": "Alcance",
    "novos_seguidores": "Novos seguidores", "posts": "Posts publicados",
}
# Cabeçalhos aceitos no posts_detalhe.csv (comparados sem acento, minúsculos, "_" = " ").
CAMPOS_CSV = {
    "url": ["url", "link", "post url", "url do post"],
    "publicado_em": ["publicado em", "data publicacao", "data hora", "data"],
    "formato": ["formato", "format", "tipo de post"],
    "tema": ["tema", "assunto", "topic"],
    "gancho": ["gancho", "hook", "primeira linha"],
    "tipo_gancho": ["tipo gancho", "tipo de gancho", "hook type"],
    "texto": ["texto", "text", "conteudo"],
    "caracteres": ["caracteres", "chars", "tamanho"],
    "impressoes": ["impressoes", "impressions"],
    "alcance": ["alcance", "membros alcancados", "members reached"],
    "reacoes": ["reacoes", "reactions"],
    "comentarios": ["comentarios", "comments"],
    "compartilhamentos": ["compartilhamentos", "reposts", "shares", "republicacoes"],
    "salvamentos": ["salvamentos", "saves"],
    "envios": ["envios", "sends"],
    "seguidores_ganhos": ["seguidores ganhos", "followers gained", "seguidores"],
    "visitas_perfil": ["visitas perfil", "visitas ao perfil", "visualizacoes do perfil",
                       "profile viewers"],
    "coletado_em": ["coletado em", "collected at"],
}


# ---------------------------------------------------------------- utilidades

def norm(s):
    s = unicodedata.normalize("NFKD", str(s if s is not None else ""))
    return s.encode("ascii", "ignore").decode().strip().lower()


def num(v):
    """Número de uma célula (o export novo grava tudo como texto: "6530", "1,234")."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "")
    s = s.replace("%", "").replace("<", "")
    if not s or s in {"-", "—", "–"}:
        return None
    if re.fullmatch(r"-?\d{1,3}([.,]\d{3})+", s):
        s = re.sub(r"[.,]", "", s)
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def to_date(v, dmy=False):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 30000 < v < 80000:  # número de série do Excel
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip()
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
    if m:
        a, b, y = int(m[1]), int(m[2]), int(m[3])
        y += 2000 if y < 100 else 0
        d, mo = (a, b) if dmy else (b, a)
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    return None


def datas_coluna(valores):
    """Converte uma coluna inteira de datas. O LinkedIn exporta M/D/AAAA; se algum
    primeiro campo passar de 12, a coluna é D/M/AAAA."""
    textos = [str(v) for v in valores if isinstance(v, str)]
    dmy = any(re.match(r"\s*(\d{1,2})/", t) and int(re.match(r"\s*(\d{1,2})/", t)[1]) > 12
              for t in textos)
    return [to_date(v, dmy) for v in valores]


def to_datetime(v, fuso):
    if isinstance(v, datetime):
        return v if v.tzinfo else v.replace(tzinfo=fuso)
    s = str(v or "").strip().replace("T", " ")
    m = re.match(r"(\d{4}-\d{2}-\d{2})[ ](\d{1,2}):(\d{2})", s)
    if m:
        d = date.fromisoformat(m[1])
        return datetime(d.year, d.month, d.day, int(m[2]), int(m[3]), tzinfo=fuso)
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})[ ]+(\d{1,2}):(\d{2})", s)
    if m:  # no CSV preenchido à mão, D/M/AAAA HH:MM
        try:
            return datetime(int(m[3]), int(m[2]), int(m[1]), int(m[4]), int(m[5]), tzinfo=fuso)
        except ValueError:
            return None
    return None


def obter_fuso(nome):
    try:
        from zoneinfo import ZoneInfo
        return ZoneInfo(nome)
    except Exception:
        if nome == "America/Sao_Paulo":  # Windows sem tzdata; Brasil sem horário de verão
            return timezone(timedelta(hours=-3), "BRT")
        return timezone.utc


def post_id(url):
    m = re.search(r"(\d{16,20})", str(url or ""))
    return m[1] if m else None


def chave_post(url):
    return post_id(url) or norm(url).split("?")[0].rstrip("/")


def momento_do_id(pid, fuso):
    """Instante da publicação codificado no ID do post (ms = ID >> 22)."""
    try:
        ms = int(pid) >> 22
    except (TypeError, ValueError):
        return None
    if not 1.2e12 < ms < 2.6e12:
        return None
    return datetime.fromtimestamp(ms / 1000, tz=timezone.utc).astimezone(fuso)


def media(xs):
    xs = [x for x in xs if x is not None]
    return statistics.mean(xs) if xs else None


def mediana(xs):
    xs = [x for x in xs if x is not None]
    return statistics.median(xs) if xs else None


def variacao(atual, base):
    if atual is None or not base:
        return None
    return atual / base - 1


def fi(x):
    return "—" if x is None else f"{x:,.0f}".replace(",", ".")


def fd(x, casas=1):
    return "—" if x is None else f"{x:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fp(x, casas=1):
    return "—" if x is None else fd(x * 100, casas) + "%"


def fv(x):
    if x is None:
        return "—"
    return ("+" if x >= 0 else "−") + fd(abs(x) * 100, 0) + "%"


def fdata(d):
    return d.strftime("%d/%m/%Y") if d else "—"


def corta(s, n=90):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


def faixa_horaria(h):
    if h is None:
        return None
    return next(r for a, b, r in FAIXAS if a <= h < b)


def tamanho(chars):
    if not chars:
        return None
    return next(r for lim, r in TAMANHOS if chars <= lim)


def classificar_gancho(g):
    """Heurística para a 1ª linha do post. A coluna tipo_gancho do CSV tem prioridade."""
    if not g:
        return None
    t = g.strip()
    n = norm(t)
    if "?" in t[:200] and (t.rstrip().endswith("?") or n.startswith(
            ("voce ", "sabia", "por que", "como ", "o que", "qual", "quando", "quem", "sera "))):
        return "Pergunta"
    if re.match(r"^\d+\s+(coisas|dicas|licoes|erros|passos|motivos|formas|sinais|perguntas|"
                r"pontos|regras|mitos|aprendizados|lições|maneiras|tendencias|ferramentas)", n) \
            or re.match(r"^\d+[.)]\s", n):
        return "Lista"
    if re.search(r"%|r\$|us\$|\bmilho|\bbilho|\btrilho|\bmil\b|\d+[.,]\d+", n) \
            or re.search(r"\b(?!19\d\d\b|20\d\d\b)\d{3,}\b", n):
        return "Dado/número"
    if n.startswith(("eu ", "ontem", "hoje eu", "ha ", "quando eu", "semana passada", "lembro",
                     "minha ", "meu ", "nunca vou", "certa vez", "no meu ", "na minha ", "fui ")):
        return "História pessoal"
    if n.startswith(("nao ", "pare ", "esqueca", "ninguem", "todo mundo", "a maioria", "o maior erro",
                     "mito", "a verdade", "opiniao impopular", "impopular", "pouca gente", "chega de",
                     "desculpe", "vou ser sincero", "unpopular")):
        return "Contraintuitivo/provocação"
    if n.startswith(("novo", "nova", "saiu", "acaba de", "acabou de", "comunicado", "anuncio",
                     "e oficial", "atencao", "urgente", "o banco central", "o bc ", "a cvm", "o coaf",
                     "foi publicad", "publicad")) or "novidade" in n:
        return "Novidade/notícia"
    if t[:1] in "\"“'«":
        return "Citação"
    return "Afirmação direta"


# ---------------------------------------------------------------- leitura do export

def localizar_export(caminho):
    p = Path(caminho).expanduser()
    if p.is_file():
        return p
    if p.is_dir():
        xs = [x for x in p.glob("*.xlsx") if not x.name.startswith("~$")]
        preferidos = [x for x in xs if x.name.startswith(("Content_", "AggregateAnalytics_"))]
        xs = preferidos or xs
        if xs:
            return max(xs, key=lambda x: x.stat().st_mtime)
    sys.exit(f"ERRO: export não encontrado em {caminho}")


def achar_cabecalho(linhas, *testes):
    """Primeira linha em que cada teste é satisfeito por alguma célula."""
    for i, r in enumerate(linhas):
        cel = [norm(c) for c in r]
        if all(any(t(c) for c in cel) for t in testes):
            return i, cel
    return None, None


def eh_data(c):
    return c in ("date", "data") or c.startswith(("date", "data "))


def ler_export(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    abas = {norm(ws.title): [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    wb.close()

    def aba(*nomes):
        for n, rows in abas.items():
            if any(n.startswith(x) for x in nomes):
                return rows
        return None

    avisos = []
    # ENGAGEMENT: Date | Impressions | Engagements (uma linha por dia)
    diario = {}
    eng = aba("engagement", "engajamento")
    candidatos = [eng] if eng else list(abas.values())
    for rows in candidatos:
        i, cel = achar_cabecalho(rows, eh_data, lambda c: "impress" in c)
        if i is None:
            continue
        cd = next(j for j, c in enumerate(cel) if eh_data(c))
        ci = next(j for j, c in enumerate(cel) if "impress" in c)
        ce = next((j for j, c in enumerate(cel) if "engag" in c or "engaj" in c), None)
        corpo = [r for r in rows[i + 1:] if len(r) > cd and r[cd] not in (None, "")]
        for r, d in zip(corpo, datas_coluna([r[cd] for r in corpo])):
            if d:
                diario[d] = {"impressoes": num(r[ci]) or 0,
                             "engajamentos": (num(r[ce]) or 0) if ce is not None else None}
        if diario:
            break
    if not diario:
        sys.exit("ERRO: não achei a aba ENGAGEMENT (Date/Impressions/Engagements) no export.")

    # FOLLOWERS: "Total followers on M/D/AAAA[:]" | N ; depois Date | New followers
    seguidores_total, seguidores_total_data = None, None
    fol = aba("followers", "seguidores")
    if fol:
        for r in fol[:5]:
            for j, c in enumerate(r):
                if norm(c).startswith("total"):
                    m = re.search(r"(\d{1,2}/\d{1,2}/\d{2,4}|\d{4}-\d{2}-\d{2})", str(c))
                    seguidores_total_data = to_date(m[1]) if m else None
                    seguidores_total = next((num(v) for v in r[j + 1:] if num(v) is not None), None)
        i, cel = achar_cabecalho(fol, eh_data, lambda c: "follow" in c or "segu" in c)
        if i is not None:
            cd = next(j for j, c in enumerate(cel) if eh_data(c))
            cf = next(j for j, c in enumerate(cel) if ("follow" in c or "segu" in c) and j != cd)
            corpo = [r for r in fol[i + 1:] if len(r) > cd and r[cd] not in (None, "")]
            for r, d in zip(corpo, datas_coluna([r[cd] for r in corpo])):
                if d:
                    diario.setdefault(d, {"impressoes": 0, "engajamentos": 0})
                    diario[d]["novos_seguidores"] = num(r[cf]) or 0
    else:
        avisos.append("O export não tem a aba FOLLOWERS; sem novos seguidores por semana.")

    # TOP POSTS: dois blocos lado a lado (URL | data | engajamentos) e (URL | data | impressões)
    posts, truncado = {}, False
    top = aba("top posts", "principais", "melhores")
    if top:
        i, cel = achar_cabecalho(top, lambda c: "url" in c or c in ("link", "link do post"))
        if i is not None:
            grupos = []
            for j, c in enumerate(cel):
                if not ("url" in c or c in ("link", "link do post")):
                    continue
                prox = range(j + 1, min(j + 4, len(cel)))
                cd = next((k for k in prox if "date" in cel[k] or "data" in cel[k]), None)
                cm = next((k for k in prox if any(x in cel[k] for x in ("engag", "engaj", "impress"))), None)
                if cm is not None:
                    grupos.append((j, cd, cm, "impressoes" if "impress" in cel[cm] else "engajamentos"))
            for j, cd, cm, campo in grupos:
                corpo = [r for r in top[i + 1:]
                         if len(r) > cm and r[j] and str(r[j]).strip().lower().startswith("http")]
                truncado |= len(corpo) >= 50
                datas = datas_coluna([r[cd] if cd is not None else None for r in corpo])
                for r, d in zip(corpo, datas):
                    url = str(r[j]).strip()
                    p = posts.setdefault(chave_post(url), {"url": url})
                    p["data_export"] = p.get("data_export") or d
                    v = num(r[cm])
                    if v is not None:
                        p[campo] = max(p.get(campo) or 0, v)
    if not posts:
        avisos.append("O export não trouxe a aba TOP POSTS (ou ela veio vazia).")
    if truncado:
        avisos.append("A lista de posts do export chegou ao limite de 50 por ranking: posts mais "
                      "fracos do ano podem estar fora, o que puxa as médias históricas por post para cima.")

    # DISCOVERY: período, impressões e membros alcançados do período todo
    descoberta = {}
    for r in aba("discovery", "descoberta") or []:
        if len(r) < 2:
            continue
        rot = norm(r[0])
        if "performance" in rot or "desempenho" in rot:
            descoberta["periodo"] = str(r[1] or "").strip()
        elif "impress" in rot:
            descoberta["impressoes"] = num(r[1])
        elif "reach" in rot or "alcanc" in rot:
            descoberta["alcance"] = num(r[1])

    return {"diario": diario, "posts": posts, "seguidores_total": seguidores_total,
            "seguidores_total_data": seguidores_total_data, "descoberta": descoberta,
            "avisos": avisos}


def ler_detalhes(caminho):
    """posts_detalhe.csv: uma linha por coleta; valores não vazios mais recentes prevalecem."""
    texto = Path(caminho).read_text(encoding="utf-8-sig")
    if not texto.strip():
        return {}
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    except csv.Error:
        dialeto = csv.excel
    leitor = csv.reader(texto.splitlines(), dialeto)
    cab = next(leitor, [])
    mapa = {}
    for j, h in enumerate(cab):
        h = norm(h).replace("_", " ")
        for campo, aliases in CAMPOS_CSV.items():
            if h in aliases and campo not in mapa.values():
                mapa[j] = campo
                break
    if "url" not in mapa.values():
        sys.exit(f"ERRO: {caminho} precisa de uma coluna 'url'.")
    det = {}
    for linha in leitor:
        reg = {mapa[j]: v.strip() for j, v in enumerate(linha) if j in mapa and v.strip()}
        if not reg.get("url"):
            continue
        d = det.setdefault(chave_post(reg["url"]), {})
        d.update(reg)
    return det


# ---------------------------------------------------------------- montagem dos posts

def montar_posts(posts_export, detalhes, fuso):
    chaves = list(posts_export) + [k for k in detalhes if k not in posts_export]
    posts = []
    for k in chaves:
        p = dict(posts_export.get(k, {}))
        d = detalhes.get(k, {})
        p["_id"] = k
        p["url"] = p.get("url") or d.get("url")
        for campo in METRICAS_DETALHE:
            v = num(d.get(campo))
            if v is None:
                continue
            # impressões: o número mais alto (o detalhe costuma ser mais recente que o export)
            p[campo] = max(v, p.get(campo) or 0) if campo == "impressoes" else v
        if p.get("engajamentos") is None and any(p.get(c) is not None for c in
                                                 ("reacoes", "comentarios", "compartilhamentos")):
            p["engajamentos"] = sum(p.get(c) or 0 for c in
                                    ("reacoes", "comentarios", "compartilhamentos", "salvamentos", "envios"))
            p["engajamento_somado"] = True
        if p.get("impressoes") and p.get("engajamentos") is not None:
            p["taxa_engajamento"] = p["engajamentos"] / p["impressoes"]

        # data e hora: ID do link > publicado_em do CSV > data do export
        quando = momento_do_id(post_id(p["url"]), fuso)
        csv_dt = to_datetime(d.get("publicado_em"), fuso) if d.get("publicado_em") else None
        data_ref = p.get("data_export") or (csv_dt.date() if csv_dt else None) \
            or (to_date(d.get("publicado_em")) if d.get("publicado_em") else None)
        if quando and data_ref and abs((quando.date() - data_ref).days) > 1:
            quando = None  # ID não bate com a data do export; fica só a data
        quando = quando or csv_dt
        p["data"] = quando.date() if quando else data_ref
        p["hora"] = quando.hour if quando else None
        p["horario"] = quando.strftime("%H:%M") if quando else None
        if not p["data"]:
            continue
        p["dia_semana"] = DIAS[p["data"].weekday()]
        p["faixa_horaria"] = faixa_horaria(p["hora"])

        texto = d.get("texto") or ""
        p["texto"] = texto
        p["gancho"] = d.get("gancho") or next((l.strip() for l in texto.splitlines() if l.strip()), "")
        p["tipo_gancho"] = d.get("tipo_gancho") or classificar_gancho(p["gancho"])
        chars = int(num(d.get("caracteres")) or 0) or len(texto)
        p["caracteres"] = chars or None
        p["tamanho"] = tamanho(chars)
        p["formato"] = (d.get("formato") or "").strip().lower() or None
        p["tema"] = (d.get("tema") or "").strip() or None
        posts.append(p)
    posts.sort(key=lambda p: (p["data"], p.get("horario") or ""), reverse=True)
    return posts


# ---------------------------------------------------------------- análises

def semanas(diario, posts, fim):
    """Blocos de 7 dias terminando em fim, fim-7, ... (só blocos completos). [0] = semana atual."""
    inicio = min(diario)
    blocos = []
    k = 0
    while True:
        b_fim = fim - timedelta(days=7 * k)
        b_ini = b_fim - timedelta(days=6)
        if b_ini < inicio:
            break
        dias = [diario.get(b_ini + timedelta(days=i), {}) for i in range(7)]
        imp = sum(x.get("impressoes") or 0 for x in dias)
        eng = sum(x.get("engajamentos") or 0 for x in dias)
        blocos.append({
            "inicio": b_ini, "fim": b_fim, "impressoes": imp, "engajamentos": eng,
            "taxa_engajamento": eng / imp if imp else None,
            "novos_seguidores": sum(x.get("novos_seguidores") or 0 for x in dias),
            "posts": sum(1 for p in posts if b_ini <= p["data"] <= b_fim),
        })
        k += 1
    return blocos


def comparar_semana(blocos):
    atual, hist = blocos[0], blocos[1:]
    linhas = []
    for campo in ("impressoes", "engajamentos", "taxa_engajamento", "novos_seguidores", "posts"):
        vals = [b[campo] for b in hist if b[campo] is not None]
        m = media(vals)
        todos = sorted([b[campo] for b in blocos if b[campo] is not None], reverse=True)
        pos = todos.index(atual[campo]) + 1 if atual[campo] is not None else None
        linhas.append({"metrica": campo, "rotulo": ROTULOS[campo], "semana": atual[campo],
                       "media_historica": m, "mediana_historica": mediana(vals),
                       "variacao": variacao(atual[campo], m),
                       "posicao": pos, "semanas": len(todos)})
    return linhas


def comparar_posts(posts, ini):
    da_semana = [p for p in posts if p["data"] >= ini]
    hist = [p for p in posts if p["data"] < ini]
    linhas = []
    for campo in METRICAS_POST:
        vs = [p.get(campo) for p in da_semana if p.get(campo) is not None]
        vh = [p.get(campo) for p in hist if p.get(campo) is not None]
        if not vs and not vh:
            continue
        linhas.append({"metrica": campo, "rotulo": ROTULOS[campo], "semana": media(vs), "n_semana": len(vs),
                       "media_historica": media(vh), "mediana_historica": mediana(vh), "n_historico": len(vh),
                       "variacao": variacao(media(vs), media(vh))})
    return da_semana, linhas


def calcular_scores(posts):
    """Score 0–100: média dos percentis do post nas métricas disponíveis para a maioria."""
    usadas = []
    for m in METRICAS_SCORE:
        vals = sorted(p[m] for p in posts if p.get(m) is not None)
        if len(vals) >= max(3, 0.5 * len(posts)):
            usadas.append((m, vals))
    for p in posts:
        pcts = []
        for m, vals in usadas:
            x = p.get(m)
            if x is None:
                continue
            menores = bisect_left(vals, x)
            iguais = bisect_right(vals, x) - menores
            pcts.append((menores + 0.5 * iguais) / len(vals))
        p["score"] = 100 * sum(pcts) / len(pcts) if pcts else None
    return [m for m, _ in usadas]


def padroes(posts, top_ids):
    med_geral = mediana([p.get("impressoes") for p in posts])
    res = []
    for chave, rotulo in DIMENSOES:
        com = [p for p in posts if p.get(chave)]
        top_com = [p for p in com if p["_id"] in top_ids]
        if len(com) < 3:
            res.append({"chave": chave, "rotulo": rotulo, "cobertura": len(com), "linhas": []})
            continue
        grupos = defaultdict(list)
        for p in com:
            grupos[p[chave]].append(p)
        linhas = []
        for valor, ps in grupos.items():
            nt = sum(1 for p in ps if p["_id"] in top_ids)
            pct_todos = len(ps) / len(com)
            pct_top = nt / len(top_com) if top_com else 0
            linhas.append({
                "valor": valor, "n": len(ps), "n_top": nt, "pct_top": pct_top, "pct_todos": pct_todos,
                "lift": pct_top / pct_todos if pct_todos else None,
                "mediana_impressoes": mediana([p.get("impressoes") for p in ps]),
                "mediana_taxa": mediana([p.get("taxa_engajamento") for p in ps]),
                "media_score": media([p.get("score") for p in ps]),
            })
        linhas.sort(key=lambda l: (-l["n_top"], -(l["lift"] or 0), -l["n"]))
        res.append({"chave": chave, "rotulo": rotulo, "cobertura": len(com), "n_top": len(top_com),
                    "mediana_geral_impressoes": med_geral, "linhas": linhas})
    return res


def aprendizados_candidatos(comp_semana, posts_semana, posts, top, pads):
    cands = []
    imp = next(l for l in comp_semana if l["metrica"] == "impressoes")
    tx = next(l for l in comp_semana if l["metrica"] == "taxa_engajamento")
    seg = next(l for l in comp_semana if l["metrica"] == "novos_seguidores")
    if imp["variacao"] is not None:
        partes = [f"Impressões da semana: {fi(imp['semana'])} ({fv(imp['variacao'])} vs a média semanal "
                  f"de {fi(imp['media_historica'])}; {imp['posicao']}ª melhor de {imp['semanas']} semanas)"]
        if tx["semana"] is not None and tx["media_historica"]:
            partes.append(f"taxa de engajamento {fp(tx['semana'])} (média {fp(tx['media_historica'])})")
        if seg["variacao"] is not None:
            partes.append(f"{fi(seg['semana'])} novos seguidores ({fv(seg['variacao'])})")
        forca = max(abs(l["variacao"] or 0) for l in (imp, tx, seg))
        cands.append({"tipo": "semana", "forca": 1 + forca, "texto": "; ".join(partes) + "."})

    score_geral = media([p.get("score") for p in posts]) or 50
    for d in pads:
        if not d["linhas"]:
            continue
        nome = d["rotulo"].lower()
        # padrão só conta se o grupo inteiro também performa acima da média (não só os do topo)
        boas = [l for l in d["linhas"] if l["n_top"] >= 2 and l["n"] >= 3 and (l["lift"] or 0) >= 1.4
                and (l["media_score"] or 0) >= score_geral + 5]
        if boas:
            l = max(boas, key=lambda l: (l["lift"] - 1) * l["n_top"])
            cands.append({"tipo": f"padrao:{d['chave']}", "forca": (l["lift"] - 1) * l["n_top"] / 2,
                          "texto": f"{d['rotulo']}: {l['n_top']} dos {d['n_top']} melhores posts "
                                   f"(com {nome} identificado) são \"{l['valor']}\", contra {fp(l['pct_todos'], 0)} "
                                   f"de todos os posts ({fd(l['lift'])}x). Mediana de "
                                   f"{fi(l['mediana_impressoes'])} impressões vs {fi(d['mediana_geral_impressoes'])} no geral."})
        med = d.get("mediana_geral_impressoes")
        fracas = [l for l in d["linhas"] if l["n"] >= 4 and l["n_top"] == 0 and med
                  and (l["mediana_impressoes"] or 0) < 0.6 * med]
        for l in fracas[:1]:
            queda = 1 - (l["mediana_impressoes"] or 0) / med
            cands.append({"tipo": f"fraco:{d['chave']}", "forca": queda * l["n"] / 6,
                          "texto": f"{d['rotulo']} \"{l['valor']}\" ({l['n']} posts) não aparece entre os melhores: "
                                   f"mediana de {fi(l['mediana_impressoes'])} impressões, {fp(queda, 0)} abaixo do geral."})

    ranqueados = [p for p in posts if p.get("score") is not None]
    if posts_semana and ranqueados:
        melhor = max((p for p in posts_semana if p.get("score") is not None), key=lambda p: p["score"], default=None)
        if melhor:
            acima = sum(1 for p in ranqueados if p["score"] < melhor["score"]) / len(ranqueados)
            cands.append({"tipo": "post_semana", "forca": abs(acima - 0.5) * 2,
                          "texto": f"Melhor post da semana ({fdata(melhor['data'])}"
                                   f"{', ' + melhor['formato'] if melhor.get('formato') else ''}"
                                   f"{', gancho: “' + corta(melhor['gancho'], 70) + '”' if melhor.get('gancho') else ''}): "
                                   f"{fi(melhor.get('impressoes'))} impressões, taxa {fp(melhor.get('taxa_engajamento'))}; "
                                   f"supera {fp(acima, 0)} dos posts dos últimos 12 meses."})
    elif not posts_semana:
        cands.append({"tipo": "sem_post", "forca": 0.8,
                      "texto": "Nenhum post publicado na semana: o resultado da semana vem de posts antigos."})

    tops = [p["caracteres"] for p in top if p.get("caracteres")]
    resto = [p["caracteres"] for p in posts if p.get("caracteres") and p not in top]
    if len(tops) >= 3 and len(resto) >= 3:
        mt, mr = mediana(tops), mediana(resto)
        if mr and abs(mt / mr - 1) >= 0.25:
            cands.append({"tipo": "tamanho_num", "forca": abs(mt / mr - 1),
                          "texto": f"Tamanho: os melhores posts têm mediana de {fi(mt)} caracteres, "
                                   f"contra {fi(mr)} nos demais ({fv(mt / mr - 1)})."})
    cands.sort(key=lambda c: -c["forca"])
    return cands


# ---------------------------------------------------------------- saída

def ler_aprendizados(caminho):
    if not caminho:
        return []
    texto = Path(caminho).read_text(encoding="utf-8").strip()
    blocos = [b.strip() for b in re.split(r"\n\s*\n|\n(?=\s*\d+[.)]\s)", texto) if b.strip()]
    if len(blocos) == 1:
        blocos = [l.strip() for l in texto.splitlines() if l.strip()]
    return [re.sub(r"^\s*(\d+[.)]|[-*•])\s*", "", b) for b in blocos]


def linha_post_md(p, i=None):
    cols = [] if i is None else [str(i)]
    cols += [fdata(p["data"]), p.get("horario") or "—", p.get("formato") or "—", corta(p.get("tema"), 30) or "—",
             corta(p.get("gancho"), 70).replace("|", "/") or "—", fi(p.get("impressoes")),
             fp(p.get("taxa_engajamento")), fi(p.get("comentarios")), fi(p.get("compartilhamentos")),
             fi(p.get("seguidores_ganhos")), fi(p.get("score")), f"[link]({p['url']})" if p.get("url") else "—"]
    return "| " + " | ".join(cols) + " |"


def escrever_md(caminho, ctx):
    ini, fim = ctx["ini"], ctx["fim"]
    L = [f"# Relatório semanal do LinkedIn: {ini:%d/%m} a {fim:%d/%m/%Y}", ""]
    L.append(f"Gerado em {datetime.now():%d/%m/%Y %H:%M} a partir de `{ctx['export'].name}`"
             + (f" (período do export: {ctx['descoberta']['periodo']})" if ctx["descoberta"].get("periodo") else "")
             + f". {ctx['n_posts']} posts no período; {ctx['n_detalhados']} com detalhes coletados no navegador.")
    L.append("")
    if ctx["aprendizados"]:
        L += ["## 3 aprendizados da semana", ""] + [f"{i}. {a}" for i, a in enumerate(ctx["aprendizados"], 1)]
    else:
        L += ["## Aprendizados candidatos (gerados automaticamente, revisar)", ""]
        L += [f"- {c['texto']}" for c in ctx["candidatos"][:6]] or ["- (sem padrões fortes nesta semana)"]
    L += ["", "## Semana x média histórica (blocos de 7 dias nos últimos 12 meses)", "",
          "| Métrica | Semana | Média semanal | Mediana semanal | Variação | Posição |",
          "|---|---:|---:|---:|---:|---:|"]
    for l in ctx["comp_semana"]:
        f = fp if l["metrica"] == "taxa_engajamento" else fi
        L.append(f"| {l['rotulo']} | {f(l['semana'])} | {f(l['media_historica'])} | {f(l['mediana_historica'])} | "
                 f"{fv(l['variacao'])} | {l['posicao'] or '—'}ª de {l['semanas']} |")
    if ctx["comp_posts"]:
        L += ["", "### Por post (posts publicados na semana x posts anteriores)", "",
              "| Métrica | Média dos posts da semana | Média histórica por post | Mediana histórica | Variação |",
              "|---|---:|---:|---:|---:|"]
        for l in ctx["comp_posts"]:
            f = fp if l["metrica"] == "taxa_engajamento" else fi
            L.append(f"| {l['rotulo']} | {f(l['semana'])} (n={l['n_semana']}) | {f(l['media_historica'])} "
                     f"(n={l['n_historico']}) | {f(l['mediana_historica'])} | {fv(l['variacao'])} |")
    cab = ("| Data | Hora | Formato | Tema | Gancho | Impressões | Taxa eng. | Coment. | Compart. | "
           "Seguidores | Score | Link |")
    sep = "|---|---|---|---|---|---:|---:|---:|---:|---:|---:|---|"
    L += ["", "## Posts da semana", ""]
    if ctx["posts_semana"]:
        L += [cab, sep] + [linha_post_md(p) for p in ctx["posts_semana"]]
        L += ["", "_Posts com menos de 7 dias ainda estão acumulando impressões; compare com cautela._"]
    else:
        L.append("Nenhum post publicado na semana.")
    L += ["", "## O que os melhores posts têm em comum", "",
          f"Melhores posts = os {len(ctx['top'])} de maior score (0–100: média dos percentis em "
          f"{', '.join(ROTULOS[m].lower() for m in ctx['metricas_score'])}).", ""]
    for d in ctx["padroes"]:
        if not d["linhas"]:
            L.append(f"- **{d['rotulo']}**: sem dado suficiente ({d['cobertura']} posts com essa informação).")
            continue
        l0 = d["linhas"][0]
        melhor = max((l for l in d["linhas"] if l["n"] >= 3), key=lambda l: l["mediana_impressoes"] or 0, default=None)
        txt = (f"- **{d['rotulo']}**: {l0['n_top']} de {d['n_top']} melhores posts são \"{l0['valor']}\" "
               f"({fp(l0['pct_top'], 0)} dos melhores vs {fp(l0['pct_todos'], 0)} de todos)")
        if melhor:
            txt += f"; maior mediana de impressões: \"{melhor['valor']}\" ({fi(melhor['mediana_impressoes'])}, n={melhor['n']})"
        if d["cobertura"] < ctx["n_posts"]:
            txt += f". Base: {d['cobertura']} de {ctx['n_posts']} posts"
        L.append(txt + ".")
    L += ["", "### Os melhores posts", "", cab.replace("| Data", "| # | Data"), "|---" + sep]
    L += [linha_post_md(p, i) for i, p in enumerate(ctx["top"], 1)]
    L += ["", "## Limitações dos dados", ""] + [f"- {a}" for a in ctx["avisos"]]
    Path(caminho).write_text("\n".join(L) + "\n", encoding="utf-8")


CAB = Font(bold=True, color="FFFFFF")
FUNDO = PatternFill("solid", fgColor="0A66C2")
TITULO = Font(bold=True, size=14)


def tabela(ws, linha, cabecalho, linhas, formatos=None, larguras=None):
    for j, h in enumerate(cabecalho, 1):
        c = ws.cell(row=linha, column=j, value=h)
        c.font, c.fill = CAB, FUNDO
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, valores in enumerate(linhas, linha + 1):
        for j, v in enumerate(valores, 1):
            c = ws.cell(row=i, column=j, value=v)
            if formatos and formatos[j - 1]:
                c.number_format = formatos[j - 1]
    for j, w in enumerate(larguras or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return linha + len(linhas) + 2


def escrever_xlsx(caminho, ctx):
    wb = openpyxl.Workbook()
    INT, PCT, VAR = "#,##0", "0.0%", "+0%;-0%;0%"

    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Relatório semanal do LinkedIn: {ctx['ini']:%d/%m} a {ctx['fim']:%d/%m/%Y}"
    ws["A1"].font = TITULO
    ws["A2"] = f"Export: {ctx['export'].name}  |  posts no período: {ctx['n_posts']}  |  com detalhes: {ctx['n_detalhados']}"
    r = 4
    ws.cell(row=r, column=1, value="3 aprendizados da semana" if ctx["aprendizados"]
            else "Aprendizados candidatos (revisar)").font = Font(bold=True, size=12)
    for i, a in enumerate(ctx["aprendizados"] or [c["texto"] for c in ctx["candidatos"][:6]], 1):
        c = ws.cell(row=r + i, column=1, value=f"{i}. {a}")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=6)
        ws.row_dimensions[r + i].height = 48
    r += max(len(ctx["aprendizados"]), min(6, len(ctx["candidatos"]))) + 2
    ws.cell(row=r, column=1, value="Semana x média semanal (12 meses)").font = Font(bold=True, size=12)
    linhas = [[l["rotulo"], l["semana"], l["media_historica"], l["mediana_historica"], l["variacao"],
               f"{l['posicao']}ª de {l['semanas']}" if l["posicao"] else "—"] for l in ctx["comp_semana"]]
    tabela(ws, r + 1, ["Métrica", "Semana", "Média semanal", "Mediana semanal", "Variação", "Posição"], linhas,
           [None, INT, INT, INT, VAR, None], [34, 16, 16, 16, 12, 14])
    for i, l in enumerate(ctx["comp_semana"], r + 2):
        if l["metrica"] == "taxa_engajamento":
            for j in (2, 3, 4):
                ws.cell(row=i, column=j).number_format = "0.00%"

    ws = wb.create_sheet("Por post")
    linhas = [[l["rotulo"], l["semana"], l["n_semana"], l["media_historica"], l["mediana_historica"],
               l["n_historico"], l["variacao"]] for l in ctx["comp_posts"]]
    tabela(ws, 1, ["Métrica", "Média posts da semana", "n semana", "Média histórica", "Mediana histórica",
                   "n histórico", "Variação"], linhas, [None, "#,##0.00", INT, "#,##0.00", "#,##0.00", INT, VAR],
           [24, 20, 10, 18, 18, 12, 12])
    for i, l in enumerate(ctx["comp_posts"], 2):
        if l["metrica"] == "taxa_engajamento":
            for j in (2, 4, 5):
                ws.cell(row=i, column=j).number_format = "0.00%"

    ws = wb.create_sheet("Posts")
    top_ids = {p["_id"] for p in ctx["top"]}
    cab = ["Data", "Dia", "Hora", "Formato", "Tema", "Tipo de gancho", "Gancho", "Caracteres", "Tamanho",
           "Impressões", "Engajamentos", "Taxa de engajamento", "Reações", "Comentários", "Compartilhamentos",
           "Salvamentos", "Envios", "Seguidores ganhos", "Visitas ao perfil", "Score", "Melhor post?",
           "Da semana?", "URL"]
    linhas = [[p["data"], p["dia_semana"], p.get("horario"), p.get("formato"), p.get("tema"), p.get("tipo_gancho"),
               corta(p.get("gancho"), 200), p.get("caracteres"), p.get("tamanho"), p.get("impressoes"),
               p.get("engajamentos"), p.get("taxa_engajamento"), p.get("reacoes"), p.get("comentarios"),
               p.get("compartilhamentos"), p.get("salvamentos"), p.get("envios"), p.get("seguidores_ganhos"),
               p.get("visitas_perfil"), round(p["score"], 1) if p.get("score") is not None else None,
               "sim" if p["_id"] in top_ids else "", "sim" if p["data"] >= ctx["ini"] else "", p.get("url")]
              for p in ctx["posts"]]
    fmts = ["DD/MM/YYYY", None, None, None, None, None, None, INT, None, INT, INT, "0.00%"] + [INT] * 7 + \
           ["0.0", None, None, None]
    tabela(ws, 1, cab, linhas, fmts, [11, 9, 7, 12, 18, 18, 50, 10, 22, 11, 12, 11, 9, 11, 13, 11, 8, 11, 10, 8,
                                      9, 9, 60])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{len(linhas) + 1}"

    ws = wb.create_sheet("Padrões")
    ws["A1"] = f"O que os {len(ctx['top'])} melhores posts têm em comum (score = média dos percentis)"
    ws["A1"].font = Font(bold=True, size=12)
    r = 3
    for d in ctx["padroes"]:
        ws.cell(row=r, column=1, value=f"{d['rotulo']} (base: {d['cobertura']} posts)").font = Font(bold=True)
        if not d["linhas"]:
            ws.cell(row=r + 1, column=1, value="Sem dado suficiente.")
            r += 3
            continue
        linhas = [[l["valor"], l["n"], l["n_top"], l["pct_todos"], l["pct_top"], l["lift"], l["mediana_impressoes"],
                   l["mediana_taxa"], round(l["media_score"], 1) if l["media_score"] is not None else None]
                  for l in d["linhas"]]
        r = tabela(ws, r + 1, ["Valor", "Posts", "Entre os melhores", "% de todos", "% dos melhores", "Lift",
                               "Mediana impressões", "Mediana taxa eng.", "Score médio"], linhas,
                   [None, INT, INT, "0%", "0%", "0.0x", INT, "0.00%", "0.0"], [30, 8, 12, 11, 13, 8, 14, 14, 11])

    ws = wb.create_sheet("Semanas")
    linhas = [[b["inicio"], b["fim"], b["impressoes"], b["engajamentos"], b["taxa_engajamento"],
               b["novos_seguidores"], b["posts"]] for b in ctx["blocos"]]
    tabela(ws, 1, ["Início", "Fim", "Impressões", "Engajamentos", "Taxa de engajamento", "Novos seguidores",
                   "Posts"], linhas, ["DD/MM/YYYY", "DD/MM/YYYY", INT, INT, "0.00%", INT, INT],
           [12, 12, 13, 13, 12, 12, 8])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Notas")
    ws.column_dimensions["A"].width = 120
    for i, a in enumerate(ctx["avisos"], 1):
        ws.cell(row=i, column=1, value=f"• {a}").alignment = Alignment(wrap_text=True)
    wb.save(caminho)


# ---------------------------------------------------------------- principal

def main():
    ap = argparse.ArgumentParser(description="Relatório semanal de conteúdo do LinkedIn.")
    ap.add_argument("--export", required=True, help="Arquivo .xlsx do LinkedIn ou pasta com os exports")
    ap.add_argument("--detalhes", help="posts_detalhe.csv com os dados coletados por post")
    ap.add_argument("--saida", default=".", help="Pasta do relatório (ex.: LinkedIn/Relatórios)")
    ap.add_argument("--aprendizados", help="Arquivo de texto com os 3 aprendizados finais")
    ap.add_argument("--fim", help="Último dia da semana analisada (AAAA-MM-DD). Padrão: último dia do export")
    ap.add_argument("--fuso", default="America/Sao_Paulo")
    a = ap.parse_args()

    fuso = obter_fuso(a.fuso)
    export = localizar_export(a.export)
    dados = ler_export(export)
    detalhes = {}
    if a.detalhes and Path(a.detalhes).expanduser().exists():
        detalhes = ler_detalhes(Path(a.detalhes).expanduser())
    elif a.detalhes:
        dados["avisos"].append(f"Arquivo de detalhes não encontrado ({a.detalhes}).")

    diario = dados["diario"]
    fim = date.fromisoformat(a.fim) if a.fim else max(diario)
    ini = fim - timedelta(days=6)
    posts = [p for p in montar_posts(dados["posts"], detalhes, fuso) if p["data"] <= fim]
    blocos = semanas(diario, posts, fim)
    if len(blocos) < 2:
        sys.exit("ERRO: o export precisa cobrir ao menos 14 dias para comparar com o histórico.")
    comp_semana = comparar_semana(blocos)
    posts_semana, comp_posts = comparar_posts(posts, ini)
    metricas_score = calcular_scores(posts)
    ranqueados = sorted([p for p in posts if p.get("score") is not None], key=lambda p: -p["score"])
    n_top = max(3, min(10, round(0.2 * len(ranqueados)))) if len(ranqueados) >= 5 else 0
    top = ranqueados[:n_top]
    pads = padroes(posts, {p["_id"] for p in top}) if top else []
    candidatos = aprendizados_candidatos(comp_semana, posts_semana, posts, top, pads)

    avisos = dados["avisos"]
    n_det = sum(1 for p in posts if p["_id"] in detalhes)
    if not detalhes:
        avisos.append("Sem posts_detalhe.csv: comentários, compartilhamentos, seguidores por post, gancho, "
                      "formato, tamanho e tema não foram analisados (o export do LinkedIn não traz esses dados).")
    elif n_det < len(posts):
        avisos.append(f"Detalhes coletados para {n_det} de {len(posts)} posts; os padrões de gancho, formato, "
                      "tamanho e tema usam só esses.")
    if any(p.get("engajamento_somado") for p in posts):
        avisos.append("Para posts fora do ranking de engajamentos do export, engajamentos = reações + comentários "
                      "+ compartilhamentos + salvamentos + envios do detalhe.")
    if not top:
        avisos.append("Menos de 5 posts com métricas: não dá para dizer o que os melhores têm em comum.")
    avisos.append(f"Semana analisada: {fdata(ini)} a {fdata(fim)} (7 dias terminando no último dia do export). "
                  "A média histórica usa os blocos de 7 dias anteriores dentro do export (até 12 meses).")
    avisos.append("Dia e horário de publicação saem do ID do post no link (fuso " + a.fuso + ").")

    saida = Path(a.saida).expanduser()
    saida.mkdir(parents=True, exist_ok=True)
    base = saida / f"Relatorio_LinkedIn_{fim.isoformat()}"
    ctx = {"ini": ini, "fim": fim, "export": export, "descoberta": dados["descoberta"], "posts": posts,
           "n_posts": len(posts), "n_detalhados": n_det, "comp_semana": comp_semana, "comp_posts": comp_posts,
           "posts_semana": posts_semana, "top": top, "metricas_score": metricas_score, "padroes": pads,
           "blocos": blocos, "candidatos": candidatos, "aprendizados": ler_aprendizados(a.aprendizados),
           "avisos": avisos}
    escrever_md(base.with_suffix(".md"), ctx)
    escrever_xlsx(base.with_suffix(".xlsx"), ctx)

    def limpo(x):
        if isinstance(x, (date, datetime)):
            return x.isoformat()
        if isinstance(x, float):
            return round(x, 4)
        if isinstance(x, dict):  # posts saem sem o texto completo, para o JSON não ficar enorme
            return {k: limpo(v) for k, v in x.items()
                    if not k.startswith("_") and not (k == "texto" and "_id" in x)}
        if isinstance(x, list):
            return [limpo(v) for v in x]
        return x

    print(json.dumps(limpo({
        "semana": {"inicio": ini, "fim": fim}, "export": str(export),
        "arquivos": [str(base.with_suffix(".xlsx")), str(base.with_suffix(".md"))],
        "posts_no_periodo": len(posts), "posts_com_detalhe": n_det,
        "seguidores_total": dados["seguidores_total"], "semana_x_historico": comp_semana,
        "posts_da_semana": posts_semana, "por_post": comp_posts,
        "melhores_posts": top, "padroes": [{**d, "linhas": d["linhas"][:4]} for d in pads],
        "aprendizados_candidatos": candidatos, "avisos": avisos,
    }), ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
